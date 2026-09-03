"""
Tests for compsai/excel.py (Module 3).

Two layers:

* Structural tests read the saved workbook back with openpyxl and check the layout the
  design contract promises (sheet names, headers, formula strings, cross-sheet references,
  fonts, number formats, chart presence). They need no external software.
* Recalculation tests hand the workbook to headless LibreOffice (tests/excel_recalc.py),
  which evaluates every formula, and then check the computed values: no Excel errors,
  EV/EBITDA matches the Python figure, the median ignores a negative-EBITDA peer, and
  the median moves when an Inputs price is edited. They skip when LibreOffice is absent.

All company data is inline and synthetic; no fixture files or network.
"""

from __future__ import annotations

import math
from datetime import date
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from compsai import excel
from compsai.excel import (
    COLUMN_SPECS,
    SHEET_COMMENTARY,
    SHEET_COMPS,
    SHEET_FOOTBALL,
    SHEET_INPUTS,
    inputs_block_start,
    write_comps_workbook,
)
from compsai.models import CommentaryResult, CompanyData
from tests.excel_recalc import recalc

AS_OF = date(2026, 9, 1)
ERROR_STRINGS = ("#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!")

# DESIGN.md section 2: extract_financials column order.
FIN_COLUMNS = [
    "fiscal_year", "period_type", "period_end", "revenue", "ebit", "da", "ebitda", "net_income",
    "eps_diluted", "total_debt", "cash", "minority_interest", "preferred", "diluted_shares",
    "total_equity", "goodwill", "intangibles", "tangible_book_value", "currency",
]
MULTIPLES_INDEX = [
    "market_cap", "ev", "ev_revenue_ttm", "ev_ebitda_ttm", "pe_ttm", "ebitda_margin",
    "net_margin", "revenue_growth_1y", "revenue_growth_3y_cagr", "p_tbv", "price", "note",
]


# --------------------------------------------------------------------------------------
# Inline company builders
# --------------------------------------------------------------------------------------


def _row(fy: int, ptype: str, end: str, revenue, ebit, da, ni, eps, debt, cash,
         shares, equity, minority=0.0, preferred=0.0, goodwill=0.0, intang=0.0) -> dict:
    ebitda = ebit + da if not (pd.isna(ebit) or pd.isna(da)) else np.nan
    return {
        "fiscal_year": fy, "period_type": ptype, "period_end": end, "revenue": revenue,
        "ebit": ebit, "da": da, "ebitda": ebitda, "net_income": ni, "eps_diluted": eps,
        "total_debt": debt, "cash": cash, "minority_interest": minority, "preferred": preferred,
        "diluted_shares": shares, "total_equity": equity, "goodwill": goodwill,
        "intangibles": intang, "tangible_book_value": equity - goodwill - intang,
        "currency": "USD",
    }


def _financials(rows: list[dict]) -> pd.DataFrame:
    return pd.DataFrame(rows, columns=FIN_COLUMNS)


def _market(ticker: str, price: float, shares: float) -> dict:
    return {
        "ticker": ticker, "price": price, "shares_outstanding": shares, "currency": "USD",
        "market_cap": price * shares, "source": "fixture", "as_of": AS_OF.isoformat(),
    }


def expected_multiples(fin: pd.DataFrame, market: dict, sector_type: str = "industrial",
                       note: str = "") -> pd.Series:
    """Multiples per DESIGN.md section 4, computed here so tests never import valuation.py."""
    fy = fin[fin["period_type"] == "FY"]
    ttm = fin[fin["period_type"] == "TTM"].iloc[-1] if (fin["period_type"] == "TTM").any() else fy.iloc[-1]
    mcap = market["price"] * market["shares_outstanding"]
    # EV = market cap + debt + minority + preferred - cash
    ev = mcap + ttm["total_debt"] + ttm["minority_interest"] + ttm["preferred"] - ttm["cash"]
    if sector_type == "bank":
        ev = np.nan
    growth_1y = fy["revenue"].iloc[-1] / fy["revenue"].iloc[-2] - 1 if len(fy) >= 2 else np.nan
    cagr = (fy["revenue"].iloc[-1] / fy["revenue"].iloc[-4]) ** (1 / 3) - 1 if len(fy) >= 4 else np.nan
    values = {
        "market_cap": mcap,
        "ev": ev,
        "ev_revenue_ttm": ev / ttm["revenue"],
        "ev_ebitda_ttm": ev / ttm["ebitda"],
        "pe_ttm": market["price"] / ttm["eps_diluted"],
        "ebitda_margin": ttm["ebitda"] / ttm["revenue"],
        "net_margin": ttm["net_income"] / ttm["revenue"],
        "revenue_growth_1y": growth_1y,
        "revenue_growth_3y_cagr": cagr,
        "p_tbv": mcap / ttm["tangible_book_value"],
        "price": market["price"],
        "note": note,
    }
    return pd.Series([values[k] for k in MULTIPLES_INDEX], index=MULTIPLES_INDEX)


def _company(ticker: str, name: str, rows: list[dict], price: float, shares: float,
             sector_type: str = "industrial", note: str = "", is_target: bool = False) -> CompanyData:
    fin = _financials(rows)
    market = _market(ticker, price, shares)
    return CompanyData(ticker=ticker, name=name, financials=fin, market=market,
                       multiples=expected_multiples(fin, market, sector_type, note),
                       is_target=is_target)


def industrial_companies() -> list[CompanyData]:
    """Four peers plus a target. BETA has negative EBITDA and a net loss (excluded from stats)."""
    alpha = _company("ALPHA", "Alpha Industries Inc.", [
        _row(2022, "FY", "2022-12-31", 1000, 180, 45, 120, 1.20, 400, 100, 100, 700, minority=10, goodwill=100, intang=50),
        _row(2023, "FY", "2023-12-31", 1100, 200, 48, 135, 1.35, 400, 100, 100, 750, minority=10, goodwill=100, intang=50),
        _row(2024, "FY", "2024-12-31", 1210, 220, 50, 145, 1.45, 400, 100, 100, 780, minority=10, goodwill=100, intang=50),
        _row(2025, "FY", "2025-12-31", 1331, 240, 50, 150, 1.50, 400, 100, 100, 800, minority=10, goodwill=100, intang=50),
        _row(2025, "TTM", "2026-06-30", 1400, 250, 50, 150, 1.50, 400, 100, 100, 800, minority=10, goodwill=100, intang=50),
    ], price=30.0, shares=100.0)
    beta = _company("BETA", "Beta Growth Corp.", [
        _row(2022, "FY", "2022-12-31", 600, -50, 15, -60, -0.30, 0, 500, 200, 900),
        _row(2023, "FY", "2023-12-31", 700, -70, 18, -75, -0.38, 0, 500, 200, 850),
        _row(2024, "FY", "2024-12-31", 800, -90, 20, -90, -0.45, 0, 500, 200, 800),
        _row(2025, "FY", "2025-12-31", 900, -100, 20, -100, -0.50, 0, 500, 200, 700),
        _row(2025, "TTM", "2026-06-30", 950, -100, 20, -100, -0.50, 0, 500, 200, 700),
    ], price=10.0, shares=200.0, note="EV/EBITDA excluded: negative EBITDA; P/E excluded: net loss")
    gamma = _company("GAMMA", "Gamma Holdings", [  # only two fiscal years -> CAGR n/a
        _row(2024, "FY", "2024-12-31", 1200, 220, 40, 120, 2.40, 300, 200, 50, 600),
        _row(2025, "FY", "2025-12-31", 1260, 230, 40, 125, 2.50, 300, 200, 50, 650),
        _row(2025, "TTM", "2026-06-30", 1300, 220, 40, 125, 2.50, 300, 200, 50, 650),
    ], price=50.0, shares=50.0)
    delta = _company("DELTA", "Delta Systems", [
        _row(2022, "FY", "2022-12-31", 500, 120, 30, 100, 0.70, 100, 300, 150, 400),
        _row(2023, "FY", "2023-12-31", 560, 130, 32, 110, 0.75, 100, 300, 150, 450),
        _row(2024, "FY", "2024-12-31", 630, 140, 35, 130, 0.85, 100, 300, 150, 500),
        _row(2025, "FY", "2025-12-31", 690, 160, 40, 145, 0.95, 100, 300, 150, 550),
        _row(2025, "TTM", "2026-06-30", 700, 160, 40, 150, 1.00, 100, 300, 150, 550),
    ], price=20.0, shares=150.0)
    target = _company("TGT", "Target Co.", [
        _row(2022, "FY", "2022-12-31", 800, 200, 40, 120, 1.50, 200, 100, 80, 500),
        _row(2023, "FY", "2023-12-31", 880, 220, 45, 130, 1.65, 200, 100, 80, 550),
        _row(2024, "FY", "2024-12-31", 960, 250, 50, 145, 1.80, 200, 100, 80, 600),
        _row(2025, "FY", "2025-12-31", 1050, 280, 50, 160, 2.00, 200, 100, 80, 650),
        _row(2025, "TTM", "2026-06-30", 1100, 280, 50, 160, 2.00, 200, 100, 80, 650),
    ], price=40.0, shares=80.0, is_target=True)
    # Deliberately NOT target-first: excel.py must reorder.
    return [alpha, beta, gamma, delta, target]


def bank_companies() -> list[CompanyData]:
    """Two banks: no operating income / D&A (EBITDA NaN), preferred stock, goodwill."""
    nan = np.nan
    bank1 = _company("BNK1", "First Fixture Bancorp", [
        _row(2022, "FY", "2022-12-31", 4000, nan, nan, 900, 3.00, 5000, 8000, 300, 9000, preferred=500, goodwill=1000, intang=200),
        _row(2023, "FY", "2023-12-31", 4300, nan, nan, 950, 3.20, 5000, 8000, 300, 9500, preferred=500, goodwill=1000, intang=200),
        _row(2024, "FY", "2024-12-31", 4600, nan, nan, 1000, 3.40, 5000, 8000, 300, 10000, preferred=500, goodwill=1000, intang=200),
        _row(2025, "FY", "2025-12-31", 4900, nan, nan, 1050, 3.50, 5000, 8000, 300, 10500, preferred=500, goodwill=1000, intang=200),
        _row(2025, "TTM", "2026-03-31", 5000, nan, nan, 1080, 3.60, 5000, 8000, 300, 10600, preferred=500, goodwill=1000, intang=200),
    ], price=45.0, shares=300.0, sector_type="bank", note="EV-based multiples are not meaningful for banks")
    bank2 = _company("BNK2", "Second Fixture Bank", [
        _row(2022, "FY", "2022-12-31", 2000, nan, nan, 400, 2.00, 2500, 4000, 200, 4000, goodwill=300),
        _row(2023, "FY", "2023-12-31", 2100, nan, nan, 420, 2.10, 2500, 4000, 200, 4200, goodwill=300),
        _row(2024, "FY", "2024-12-31", 2200, nan, nan, 440, 2.20, 2500, 4000, 200, 4400, goodwill=300),
        _row(2025, "FY", "2025-12-31", 2300, nan, nan, 460, 2.30, 2500, 4000, 200, 4600, goodwill=300),
        _row(2025, "TTM", "2026-03-31", 2350, nan, nan, 470, 2.35, 2500, 4000, 200, 4650, goodwill=300),
    ], price=30.0, shares=200.0, sector_type="bank", note="EV-based multiples are not meaningful for banks")
    return [bank1, bank2]


def sample_commentary() -> dict[str, CommentaryResult]:
    return {
        "ALPHA": CommentaryResult(
            ticker="ALPHA",
            normalization_items=[{
                "description": "Restructuring charge", "amount_usd_m": 25.0, "fiscal_year": 2025,
                "direction": "add_back", "source_quote": "we recorded restructuring charges of $25 million",
                "verified": True, "quote_word_count": 8,
            }],
            premium_discount={
                "growth_outlook": "High-single-digit growth", "margin_trajectory": "Expanding",
                "key_risks": ["Customer concentration", "FX"], "premium_or_discount": "premium",
                "rationale": "Faster growth than peers.",
            },
            source_form="10-K", source_url="https://www.sec.gov/x", filing_date="2026-02-15",
            model="claude-sonnet-4-6",
        ),
        "BETA": CommentaryResult(ticker="BETA", errors=["Claude call failed: 429"]),
    }


# --------------------------------------------------------------------------------------
# Workbook helpers
# --------------------------------------------------------------------------------------


def _header_cols(ws, header_row: int = excel.COMPS_HEADER_ROW) -> dict[str, int]:
    """Map header text -> column index for a sheet (looks columns up by header text)."""
    cols = {}
    for cell in ws[header_row]:
        if cell.value:
            cols[str(cell.value)] = cell.column
    return cols


def _company_rows(ws) -> dict[str, int]:
    """Ticker -> row on the Comps sheet."""
    ticker_col = _header_cols(ws)["Ticker"]
    rows = {}
    r = excel.COMPS_FIRST_COMPANY_ROW
    while ws.cell(row=r, column=ticker_col).value:
        rows[str(ws.cell(row=r, column=ticker_col).value)] = r
        r += 1
    return rows


def _stat_rows(ws) -> dict[str, int]:
    rows = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if label in ("Mean", "Median", "25th Percentile", "75th Percentile"):
            rows[str(label)] = r
    return rows


def _all_values(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    yield ws.title, cell.coordinate, cell.value


@pytest.fixture(scope="module")
def industrial_path(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("xlsx")
    return write_comps_workbook("test_peers", industrial_companies(), sector_type="industrial",
                                target="TGT", commentary=sample_commentary(), out_dir=out, as_of=AS_OF)


@pytest.fixture(scope="module")
def industrial_wb(industrial_path):
    return openpyxl.load_workbook(industrial_path)


@pytest.fixture(scope="module")
def bank_path(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("xlsx_bank")
    return write_comps_workbook("banks", bank_companies(), sector_type="bank", target=None,
                                commentary=None, out_dir=out, as_of=AS_OF)


@pytest.fixture(scope="module")
def bank_wb(bank_path):
    return openpyxl.load_workbook(bank_path)


# --------------------------------------------------------------------------------------
# Structural tests (no LibreOffice needed)
# --------------------------------------------------------------------------------------


def test_file_name_and_sheets(industrial_path, industrial_wb, bank_path, bank_wb):
    assert industrial_path.name == "comps_test_peers_2026-09-01.xlsx"
    assert industrial_wb.sheetnames == [SHEET_COMPS, SHEET_INPUTS, SHEET_FOOTBALL, SHEET_COMMENTARY]
    # No target -> no Football Field sheet
    assert bank_path.name == "comps_banks_2026-09-01.xlsx"
    assert bank_wb.sheetnames == [SHEET_COMPS, SHEET_INPUTS, SHEET_COMMENTARY]


def test_column_specs_export():
    assert isinstance(COLUMN_SPECS, list)
    for spec in COLUMN_SPECS:
        assert {"key", "header", "kind"} <= set(spec)
    industrial = [s["header"] for s in excel.column_specs("industrial")]
    assert industrial == [
        "Company", "Ticker", "Price", "Market Cap", "Enterprise Value", "EV / Revenue (TTM)",
        "EV / EBITDA (TTM)", "P / E (TTM)", "EBITDA Margin", "Net Margin",
        "Revenue Growth (1Y)", "Revenue CAGR (3Y)", "Notes",
    ]
    bank = [s["header"] for s in excel.column_specs("bank")]
    assert bank == [
        "Company", "Ticker", "Price", "Market Cap", "P / E (TTM)", "P / TBV", "Net Margin",
        "Revenue Growth (1Y)", "Revenue CAGR (3Y)", "Notes",
    ]
    with pytest.raises(ValueError):
        excel.column_specs("insurance")


def test_comps_title_rows_and_headers(industrial_wb):
    ws = industrial_wb[SHEET_COMPS]
    assert ws["A1"].value == "Comparable Companies Analysis — test_peers"
    assert ws["A2"].value.startswith("As of 2026-09-01 · $ in millions except per share")
    headers = [c.value for c in ws[excel.COMPS_HEADER_ROW] if c.value]
    visible = headers[:13]
    assert visible == [s["header"] for s in excel.column_specs("industrial")]
    # helper columns sit to the right of Notes, one per sector multiple
    assert headers[13:] == [
        "EV / Revenue (TTM) (stat-eligible)", "EV / EBITDA (TTM) (stat-eligible)",
        "P / E (TTM) (stat-eligible)",
    ]
    assert ws.freeze_panes == "A5"


def test_target_first_and_marked(industrial_wb):
    ws = industrial_wb[SHEET_COMPS]
    rows = _company_rows(ws)
    assert list(rows) == ["TGT", "ALPHA", "BETA", "GAMMA", "DELTA"]
    notes_col = _header_cols(ws)["Notes"]
    assert "(target — excluded from peer stats)" in ws.cell(row=rows["TGT"], column=notes_col).value
    beta_note = ws.cell(row=rows["BETA"], column=notes_col).value
    assert beta_note.startswith("EV/EBITDA excluded")
    assert ws.cell(row=rows["ALPHA"], column=notes_col).value is None  # clean -> blank, not ""


def test_every_numeric_comps_cell_is_a_formula_into_inputs(industrial_wb):
    ws = industrial_wb[SHEET_COMPS]
    cols = _header_cols(ws)
    numeric_headers = [s["header"] for s in excel.column_specs("industrial")
                       if s["kind"] not in ("text", "notes")]
    for ticker, r in _company_rows(ws).items():
        for header in numeric_headers:
            v = ws.cell(row=r, column=cols[header]).value
            assert isinstance(v, str) and v.startswith("="), (ticker, header, v)
            assert "Inputs!" in v, (ticker, header, v)


def test_comps_formulas_match_design(industrial_wb):
    ws = industrial_wb[SHEET_COMPS]
    cols = _header_cols(ws)
    rows = _company_rows(ws)
    r = rows["ALPHA"]  # second block: b = 3 + 9*1 = 12, t = 19, y = 18
    b, t, y = inputs_block_start(1), inputs_block_start(1) + 7, inputs_block_start(1) + 6
    assert (b, t, y) == (12, 19, 18)
    assert ws.cell(row=r, column=cols["Price"]).value == f"=Inputs!$B${b + 1}"
    assert ws.cell(row=r, column=cols["Market Cap"]).value == f"=Inputs!$B${b + 1}*Inputs!$D${b + 1}"
    ev = ws.cell(row=r, column=cols["Enterprise Value"]).value
    assert ev == f"=Inputs!$B${b + 1}*Inputs!$D${b + 1}+Inputs!$J${t}+Inputs!$L${t}+Inputs!$M${t}-Inputs!$K${t}"
    ev_ebitda = ws.cell(row=r, column=cols["EV / EBITDA (TTM)"]).value
    assert ev_ebitda.startswith("=IFERROR((") and ev_ebitda.endswith(f')/Inputs!$G${t},"n/a")')
    assert ws.cell(row=r, column=cols["P / E (TTM)"]).value == f'=IFERROR(Inputs!$B${b + 1}/Inputs!$I${t},"n/a")'
    assert ws.cell(row=r, column=cols["EBITDA Margin"]).value == f'=IFERROR(Inputs!$G${t}/Inputs!$D${t},"n/a")'
    assert ws.cell(row=r, column=cols["Net Margin"]).value == f'=IFERROR(Inputs!$H${t}/Inputs!$D${t},"n/a")'
    assert ws.cell(row=r, column=cols["Revenue Growth (1Y)"]).value == f'=IFERROR(Inputs!$D${y}/Inputs!$D${b + 5}-1,"n/a")'
    assert ws.cell(row=r, column=cols["Revenue CAGR (3Y)"]).value == f'=IFERROR((Inputs!$D${y}/Inputs!$D${b + 3})^(1/3)-1,"n/a")'


def test_stats_rows_are_live_formulas_over_helper_columns(industrial_wb):
    ws = industrial_wb[SHEET_COMPS]
    cols = _header_cols(ws)
    rows = _company_rows(ws)
    first, last = min(rows.values()), max(rows.values())
    stats = _stat_rows(ws)
    assert stats["Mean"] == last + 2  # one blank row after the companies
    assert [stats[k] for k in ("Mean", "Median", "25th Percentile", "75th Percentile")] == \
        [last + 2, last + 3, last + 4, last + 5]
    helper = cols["EV / EBITDA (TTM) (stat-eligible)"]
    hl = openpyxl.utils.get_column_letter(helper)
    mult_col = cols["EV / EBITDA (TTM)"]
    assert ws.cell(row=stats["Mean"], column=mult_col).value == f"=AVERAGE({hl}{first}:{hl}{last})"
    median = ws.cell(row=stats["Median"], column=mult_col).value
    assert median.startswith("=MEDIAN(") and median == f"=MEDIAN({hl}{first}:{hl}{last})"
    assert ws.cell(row=stats["25th Percentile"], column=mult_col).value == f"=QUARTILE({hl}{first}:{hl}{last},1)"
    assert ws.cell(row=stats["75th Percentile"], column=mult_col).value == f"=QUARTILE({hl}{first}:{hl}{last},3)"
    # stats only for the sector's multiples: margins/growth stay blank
    for header in ("EBITDA Margin", "Net Margin", "Revenue Growth (1Y)", "Market Cap"):
        assert ws.cell(row=stats["Median"], column=cols[header]).value is None
    # stats fonts are black formulas
    assert ws.cell(row=stats["Median"], column=mult_col).font.color.rgb.endswith("000000")


def test_helper_columns_mirror_bounds_and_blank_target(industrial_wb):
    ws = industrial_wb[SHEET_COMPS]
    cols = _header_cols(ws)
    rows = _company_rows(ws)
    g = openpyxl.utils.get_column_letter(cols["EV / EBITDA (TTM)"])
    h = openpyxl.utils.get_column_letter(cols["P / E (TTM)"])
    r = rows["ALPHA"]
    assert ws.cell(row=r, column=cols["EV / EBITDA (TTM) (stat-eligible)"]).value == \
        f'=IF(AND(ISNUMBER({g}{r}),{g}{r}>0,{g}{r}<=100),{g}{r},"")'
    assert ws.cell(row=r, column=cols["P / E (TTM) (stat-eligible)"]).value == \
        f'=IF(AND(ISNUMBER({h}{r}),{h}{r}>0,{h}{r}<=200),{h}{r},"")'
    # target helper cells evaluate to "" so the target never enters the statistics
    rt = rows["TGT"]
    for header in ("EV / Revenue (TTM) (stat-eligible)", "EV / EBITDA (TTM) (stat-eligible)",
                   "P / E (TTM) (stat-eligible)"):
        assert ws.cell(row=rt, column=cols[header]).value == '=""'
    # helper header is grey italic and narrow
    hc = ws.cell(row=excel.COMPS_HEADER_ROW, column=cols["EV / EBITDA (TTM) (stat-eligible)"])
    assert hc.font.italic and hc.font.color.rgb.endswith("808080")
    assert ws.column_dimensions[openpyxl.utils.get_column_letter(hc.column)].width <= 8


def test_inputs_block_layout(industrial_wb):
    ws = industrial_wb[SHEET_INPUTS]
    # Order on Inputs follows the Comps order: TGT (0), ALPHA (1), BETA (2), GAMMA (3), DELTA (4)
    b = inputs_block_start(1)
    assert ws.cell(row=b, column=1).value == "ALPHA" and ws.cell(row=b, column=1).font.bold
    assert ws.cell(row=b, column=2).value == "Alpha Industries Inc."
    assert ws.cell(row=b + 1, column=1).value == "Price ($)"
    assert ws.cell(row=b + 1, column=2).value == 30.0
    assert ws.cell(row=b + 1, column=3).value == "Shares out (mm)"
    assert ws.cell(row=b + 1, column=4).value == 100.0
    assert ws.cell(row=b + 1, column=5).value == "Currency" and ws.cell(row=b + 1, column=6).value == "USD"
    assert ws.cell(row=b + 1, column=7).value == "Price as of" and ws.cell(row=b + 1, column=8).value == "2026-09-01"
    headers = [ws.cell(row=b + 2, column=c).value for c in range(1, 19)]
    assert headers == [
        "Period", "Fiscal year", "Period end", "Revenue", "EBIT", "D&A", "EBITDA", "Net income",
        "Diluted EPS", "Total debt", "Cash", "Minority interest", "Preferred", "Diluted shares (mm)",
        "Total equity", "Goodwill", "Intangibles", "Tangible book value",
    ]
    # FY rows oldest at b+3 ... newest at b+6, TTM at b+7, b+8 blank
    assert [ws.cell(row=b + k, column=2).value for k in range(3, 7)] == [2022, 2023, 2024, 2025]
    assert ws.cell(row=b + 6, column=1).value == "FY" and ws.cell(row=b + 6, column=4).value == 1331
    assert ws.cell(row=b + 7, column=1).value == "TTM" and ws.cell(row=b + 7, column=4).value == 1400
    assert ws.cell(row=b + 7, column=3).value == "2026-06-30"
    assert all(ws.cell(row=b + 8, column=c).value is None for c in range(1, 19))
    # EBITDA and TBV are formulas (black); the rest are blue hardcodes
    r = b + 7
    assert ws.cell(row=r, column=7).value == f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),E{r}+F{r},"")'
    assert ws.cell(row=r, column=18).value == f'=IF(ISNUMBER(O{r}),O{r}-P{r}-Q{r},"")'
    assert ws.cell(row=b + 7, column=7).font.color.rgb.endswith("000000")
    for col in (4, 5, 6, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17):
        cell = ws.cell(row=b + 7, column=col)
        assert cell.font.color.rgb.endswith("0000FF"), (col, cell.value)
    assert ws.cell(row=b + 1, column=2).font.color.rgb.endswith("0000FF")  # price
    assert ws.cell(row=b + 1, column=4).font.color.rgb.endswith("0000FF")  # shares
    # Next block starts exactly 9 rows later
    assert ws.cell(row=inputs_block_start(2), column=1).value == "BETA"


def test_inputs_bottom_aligns_short_history(industrial_wb):
    ws = industrial_wb[SHEET_INPUTS]
    b = inputs_block_start(3)  # GAMMA has only FY2024 and FY2025
    assert ws.cell(row=b, column=1).value == "GAMMA"
    assert ws.cell(row=b + 3, column=2).value is None and ws.cell(row=b + 4, column=2).value is None
    assert ws.cell(row=b + 3, column=7).value is None  # no stray EBITDA formula on empty years
    assert ws.cell(row=b + 5, column=2).value == 2024
    assert ws.cell(row=b + 6, column=2).value == 2025  # newest always at b+6
    assert ws.cell(row=b + 7, column=1).value == "TTM"


def test_fonts_and_number_formats(industrial_wb):
    ws = industrial_wb[SHEET_COMPS]
    cols = _header_cols(ws)
    r = _company_rows(ws)["ALPHA"]
    for sheet in industrial_wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    assert cell.font.name == "Arial" and cell.font.size == 10, (sheet.title, cell.coordinate)
    assert ws.cell(row=r, column=cols["Market Cap"]).number_format == "$#,##0"
    assert ws.cell(row=r, column=cols["Price"]).number_format == "$#,##0.00"
    assert ws.cell(row=r, column=cols["EV / EBITDA (TTM)"]).number_format == '0.0"x"'
    assert ws.cell(row=r, column=cols["Net Margin"]).number_format == "0.0%"
    # formulas on Comps are black
    assert ws.cell(row=r, column=cols["EV / EBITDA (TTM)"]).font.color.rgb.endswith("000000")


def test_nan_values_become_blank_cells(bank_wb):
    ws = bank_wb[SHEET_INPUTS]
    b = inputs_block_start(0)
    assert ws.cell(row=b + 7, column=5).value is None  # EBIT NaN -> blank
    assert ws.cell(row=b + 7, column=6).value is None  # D&A NaN -> blank
    for _sheet, coord, v in _all_values(bank_wb):
        assert not (isinstance(v, str) and v.strip().lower() in ("nan", "none")), coord


def test_bank_columns_and_formulas(bank_wb):
    ws = bank_wb[SHEET_COMPS]
    headers = [c.value for c in ws[excel.COMPS_HEADER_ROW] if c.value]
    assert headers[:10] == [s["header"] for s in excel.column_specs("bank")]
    assert headers[10:] == ["P / E (TTM) (stat-eligible)", "P / TBV (stat-eligible)"]
    assert "Enterprise Value" not in headers and "EV / EBITDA (TTM)" not in headers
    cols = _header_cols(ws)
    rows = _company_rows(ws)
    r = rows["BNK1"]
    b, t = inputs_block_start(0), inputs_block_start(0) + 7
    assert ws.cell(row=r, column=cols["P / TBV"]).value == \
        f'=IFERROR((Inputs!$B${b + 1}*Inputs!$D${b + 1})/Inputs!$R${t},"n/a")'
    ptbv = openpyxl.utils.get_column_letter(cols["P / TBV"])
    assert ws.cell(row=r, column=cols["P / TBV (stat-eligible)"]).value == \
        f'=IF(AND(ISNUMBER({ptbv}{r}),{ptbv}{r}>0,{ptbv}{r}<=20),{ptbv}{r},"")'
    stats = _stat_rows(ws)
    assert ws.cell(row=stats["Median"], column=cols["P / TBV"]).value.startswith("=MEDIAN(")
    assert ws.cell(row=stats["Median"], column=cols["P / E (TTM)"]).value.startswith("=MEDIAN(")


def test_football_field_structure(industrial_wb):
    ws = industrial_wb[SHEET_FOOTBALL]
    comps = industrial_wb[SHEET_COMPS]
    assert ws["A1"].value == "Football Field — TGT (Target Co.)"
    assert ws["A2"].value == "Current price" and ws["B2"].value == "=Inputs!$B$4"  # target block b=3
    headers = [ws.cell(row=excel.FF_HEADER_ROW, column=c).value for c in range(1, 12)]
    assert headers == [
        "Method", "Target metric", "Metric value", "Low multiple (25th)", "High multiple (75th)",
        "Implied EV low", "Implied EV high", "Implied equity low", "Implied equity high",
        "Implied price low", "Implied price high",
    ]
    methods = [ws.cell(row=excel.FF_FIRST_ROW + j, column=1).value for j in range(3)]
    assert methods == ["EV / Revenue (TTM)", "EV / EBITDA (TTM)", "P / E (TTM)"]
    stats = _stat_rows(comps)
    ccols = _header_cols(comps)
    g = openpyxl.utils.get_column_letter(ccols["EV / EBITDA (TTM)"])
    r = excel.FF_FIRST_ROW + 1  # EV/EBITDA row
    assert ws.cell(row=r, column=3).value == "=Inputs!$G$10"  # TTM EBITDA of block 0 (t = 3+7)
    assert ws.cell(row=r, column=4).value == f"='Comps'!${g}${stats['25th Percentile']}"
    assert ws.cell(row=r, column=5).value == f"='Comps'!${g}${stats['75th Percentile']}"
    assert ws.cell(row=r, column=6).value == f'=IFERROR(D{r}*C{r},"n/a")'
    assert "Inputs!$J$10" in ws.cell(row=r, column=8).value  # equity = EV - debt ... + cash
    assert ws.cell(row=r, column=10).value == f'=IFERROR(H{r}/Inputs!$D$4,"n/a")'
    # P/E row: price = multiple x EPS
    r_pe = excel.FF_FIRST_ROW + 2
    assert ws.cell(row=r_pe, column=3).value == "=Inputs!$I$10"
    assert ws.cell(row=r_pe, column=10).value == f'=IFERROR(D{r_pe}*C{r_pe},"n/a")'
    # chart: horizontal stacked bar, invisible low series + (high - low) series
    assert len(ws._charts) == 1
    chart = ws._charts[0]
    assert chart.type == "bar" and chart.grouping == "stacked" and chart.overlap == 100
    assert len(chart.series) == 2
    assert chart.series[0].graphicalProperties.noFill is True
    assert "Implied share price (TGT)" in _chart_title_text(chart)


def _chart_title_text(chart) -> str:
    parts = []
    for p in chart.title.tx.rich.p:
        for run in p.r or []:
            parts.append(run.t)
    return "".join(parts)


def test_commentary_sheet(industrial_wb, bank_wb):
    ws = industrial_wb[SHEET_COMMENTARY]
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "ALPHA — Alpha Industries Inc. — premium" in texts
    assert "Growth outlook" in texts and "High-single-digit growth" in texts
    assert "Margin trajectory" in texts and "Key risks" in texts and "Rationale" in texts
    assert "Customer concentration" in texts and "FX" in texts
    for header in ("Description", "Amount ($mm)", "Fiscal year", "Direction", "Source quote",
                   "Verified in filing"):
        assert header in texts
    assert "Restructuring charge" in texts and "Yes" in texts
    assert "Source: 10-K filed 2026-02-15 — https://www.sec.gov/x" in texts
    # BETA had an error and no content; GAMMA/DELTA/TGT had no entry at all
    assert "Commentary not generated: Claude call failed: 429" in texts
    assert texts.count("Commentary not generated: no commentary result for this company") == 3
    # amounts are blue hardcodes with $ format
    for row in ws.iter_rows():
        for c in row:
            if c.value == 25.0:
                assert c.font.color.rgb.endswith("0000FF") and c.number_format == "$#,##0"
    # commentary=None -> one placeholder line per company
    ws_b = bank_wb[SHEET_COMMENTARY]
    texts_b = [str(c.value) for row in ws_b.iter_rows() for c in row if c.value is not None]
    assert texts_b.count("Commentary not generated: AI commentary was not run for this analysis") == 2


def test_ttm_missing_links_to_newest_fy(tmp_path):
    """No TTM row -> the TTM slot is a formula link to the newest FY row, not a copy."""
    rows = [
        _row(2023, "FY", "2023-12-31", 100, 20, 5, 10, 1.0, 10, 5, 10, 50),
        _row(2024, "FY", "2024-12-31", 110, 22, 5, 11, 1.1, 10, 5, 10, 55),
    ]
    c = _company("NOTTM", "No TTM Co.", rows, price=10.0, shares=10.0)
    path = write_comps_workbook("solo", [c], out_dir=tmp_path, as_of=AS_OF)
    ws = openpyxl.load_workbook(path)[SHEET_INPUTS]
    b = inputs_block_start(0)
    assert ws.cell(row=b + 7, column=1).value == "TTM (=FY)"
    assert ws.cell(row=b + 7, column=4).value == f"=D{b + 6}"
    assert ws.cell(row=b + 7, column=4).font.color.rgb.endswith("000000")


def test_target_not_in_companies_skips_football_field(tmp_path):
    comps = industrial_companies()[:2]
    path = write_comps_workbook("x", comps, target="ZZZ", out_dir=tmp_path, as_of=AS_OF)
    wb = openpyxl.load_workbook(path)
    assert SHEET_FOOTBALL not in wb.sheetnames


def test_missing_price_company_excluded_from_stats(tmp_path):
    """A blank Inputs price reads as 0 in Excel (EV = net debt), so a company with no
    market price must not feed the helper columns: Excel would otherwise average a
    bogus in-bounds multiple that valuation.py reports as NaN."""
    comps = industrial_companies()[:2]  # ALPHA, BETA
    rows = [
        _row(2024, "FY", "2024-12-31", 1000, 80, 20, 50, 0.5, 500, 100, 100, 500),
        _row(2025, "FY", "2025-12-31", 1000, 80, 20, 50, 0.5, 500, 100, 100, 500),
        _row(2025, "TTM", "2026-06-30", 1000, 80, 20, 50, 0.5, 500, 100, 100, 500),
    ]
    nopx = _company("NOPX", "No Price Co.", rows, price=np.nan, shares=100.0)
    nopx.market["market_cap"] = np.nan
    path = write_comps_workbook("noprice", comps + [nopx], out_dir=tmp_path, as_of=AS_OF)
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET_COMPS]
    cols = _header_cols(ws)
    r = _company_rows(ws)["NOPX"]
    for header in ("EV / Revenue (TTM) (stat-eligible)", "EV / EBITDA (TTM) (stat-eligible)",
                   "P / E (TTM) (stat-eligible)"):
        assert ws.cell(row=r, column=cols[header]).value == '=""'
    assert "excluded from peer stats" in ws.cell(row=r, column=cols["Notes"]).value
    # the visible cells stay contract formulas into Inputs; only the helper is blanked
    assert ws.cell(row=r, column=cols["Price"]).value == f"=Inputs!$B${inputs_block_start(2) + 1}"
    assert wb[SHEET_INPUTS].cell(row=inputs_block_start(2) + 1, column=2).value is None
    # ALPHA (priced) keeps a live helper formula
    ra = _company_rows(ws)["ALPHA"]
    assert ws.cell(row=ra, column=cols["EV / EBITDA (TTM) (stat-eligible)"]).value.startswith("=IF(AND(ISNUMBER(")


def test_commentary_key_risks_as_string_is_one_row(tmp_path):
    """A fixture or loosely parsed reply may carry key_risks as a bare string; it must
    become one 'Key risks' line, not one line per character."""
    result = CommentaryResult(
        ticker="ALPHA",
        premium_discount={"growth_outlook": "Flat", "key_risks": "Customer concentration",
                          "premium_or_discount": "inline"},
    )
    path = write_comps_workbook("risks", industrial_companies()[:1], commentary={"ALPHA": result},
                                out_dir=tmp_path, as_of=AS_OF)
    ws = openpyxl.load_workbook(path)[SHEET_COMMENTARY]
    labels = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    values = [ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)]
    assert labels.count("Key risks") == 1
    assert values.count("Customer concentration") == 1
    assert "C" not in values  # no per-character rows


def test_default_out_dir_and_today(monkeypatch, tmp_path):
    monkeypatch.setattr(excel.compsai, "OUTPUT_DIR", tmp_path / "out")
    path = write_comps_workbook("dflt", industrial_companies()[:1])
    assert path.parent == tmp_path / "out"
    assert path.name == f"comps_dflt_{date.today().isoformat()}.xlsx"


# --------------------------------------------------------------------------------------
# Recalculation tests (LibreOffice evaluates the formulas)
# --------------------------------------------------------------------------------------


def _recalc_or_skip(path: Path):
    out = recalc(path)
    if out is None:
        pytest.skip("LibreOffice (soffice) is not available or failed; formula values not verified")
    return openpyxl.load_workbook(out, data_only=True)


def _assert_no_excel_errors(wb) -> None:
    for sheet, coord, v in _all_values(wb):
        if isinstance(v, str):
            assert not any(e in v for e in ERROR_STRINGS), (sheet, coord, v)


def _eligible(values: list[float], lower: float, upper: float) -> list[float]:
    return [v for v in values if isinstance(v, (int, float)) and math.isfinite(v) and lower < v <= upper]


@pytest.fixture(scope="module")
def industrial_calc(industrial_path):
    return _recalc_or_skip(industrial_path)


def test_recalc_no_errors(industrial_calc):
    _assert_no_excel_errors(industrial_calc)


def test_recalc_ev_ebitda_matches_python(industrial_calc):
    ws = industrial_calc[SHEET_COMPS]
    cols = _header_cols(ws)
    rows = _company_rows(ws)
    by_ticker = {c.ticker: c for c in industrial_companies()}
    for ticker, r in rows.items():
        expected = by_ticker[ticker].multiples
        for header, key in (("EV / EBITDA (TTM)", "ev_ebitda_ttm"), ("P / E (TTM)", "pe_ttm"),
                            ("Enterprise Value", "ev"), ("Net Margin", "net_margin"),
                            ("Revenue Growth (1Y)", "revenue_growth_1y")):
            got = ws.cell(row=r, column=cols[header]).value
            assert abs(got - expected[key]) < 1e-6, (ticker, header, got, expected[key])
    # GAMMA has only 2 fiscal years -> CAGR is "n/a", not an error
    assert ws.cell(row=rows["GAMMA"], column=cols["Revenue CAGR (3Y)"]).value == "n/a"


def test_recalc_median_excludes_negative_ebitda_and_target(industrial_calc):
    ws = industrial_calc[SHEET_COMPS]
    cols = _header_cols(ws)
    stats = _stat_rows(ws)
    peers = [c for c in industrial_companies() if not c.is_target]
    all_vals = [float(c.multiples["ev_ebitda_ttm"]) for c in peers]
    eligible = _eligible(all_vals, 0, 100)
    assert len(eligible) == 3  # BETA (negative EBITDA) dropped
    median_cell = ws.cell(row=stats["Median"], column=cols["EV / EBITDA (TTM)"]).value
    assert abs(median_cell - float(np.median(eligible))) < 1e-6
    # including BETA or the target would give a different median
    assert abs(median_cell - float(np.median(all_vals))) > 1e-3
    assert abs(ws.cell(row=stats["Mean"], column=cols["EV / EBITDA (TTM)"]).value - float(np.mean(eligible))) < 1e-6
    assert abs(ws.cell(row=stats["25th Percentile"], column=cols["EV / EBITDA (TTM)"]).value
               - float(np.percentile(eligible, 25))) < 1e-6
    assert abs(ws.cell(row=stats["75th Percentile"], column=cols["EV / EBITDA (TTM)"]).value
               - float(np.percentile(eligible, 75))) < 1e-6
    # target helper cells evaluate to empty text
    rt = _company_rows(ws)["TGT"]
    assert ws.cell(row=rt, column=cols["EV / EBITDA (TTM) (stat-eligible)"]).value in (None, "")


def test_recalc_football_field_values(industrial_calc):
    ws = industrial_calc[SHEET_FOOTBALL]
    comps = industrial_calc[SHEET_COMPS]
    cols = _header_cols(comps)
    stats = _stat_rows(comps)
    p25 = comps.cell(row=stats["25th Percentile"], column=cols["EV / EBITDA (TTM)"]).value
    target = industrial_companies()[-1]
    ttm = target.financials.iloc[-1]
    r = excel.FF_FIRST_ROW + 1  # EV/EBITDA row
    assert ws.cell(row=r, column=3).value == ttm["ebitda"]
    assert abs(ws.cell(row=r, column=4).value - p25) < 1e-9
    implied_ev_low = p25 * ttm["ebitda"]
    implied_eq_low = implied_ev_low - ttm["total_debt"] - ttm["minority_interest"] - ttm["preferred"] + ttm["cash"]
    assert abs(ws.cell(row=r, column=6).value - implied_ev_low) < 1e-6
    assert abs(ws.cell(row=r, column=10).value - implied_eq_low / target.market["shares_outstanding"]) < 1e-6
    assert ws["B2"].value == 40.0
    r_pe = excel.FF_FIRST_ROW + 2
    p25_pe = comps.cell(row=stats["25th Percentile"], column=cols["P / E (TTM)"]).value
    assert abs(ws.cell(row=r_pe, column=10).value - p25_pe * ttm["eps_diluted"]) < 1e-6


def test_recalc_median_moves_when_input_changes(industrial_path, industrial_calc, tmp_path):
    """Edit ALPHA's price on Inputs, recalculate, and the EV/EBITDA median must change."""
    comps = industrial_calc[SHEET_COMPS]
    cols = _header_cols(comps)
    stats = _stat_rows(comps)
    before = comps.cell(row=stats["Median"], column=cols["EV / EBITDA (TTM)"]).value

    wb = openpyxl.load_workbook(industrial_path)
    ws = wb[SHEET_INPUTS]
    b = inputs_block_start(1)  # ALPHA
    assert ws.cell(row=b, column=1).value == "ALPHA"
    ws.cell(row=b + 1, column=2).value = 60.0  # price 30 -> 60
    edited = tmp_path / "edited.xlsx"
    wb.save(edited)

    after_wb = _recalc_or_skip(edited)
    after_ws = after_wb[SHEET_COMPS]
    after = after_ws.cell(row=stats["Median"], column=cols["EV / EBITDA (TTM)"]).value
    assert abs(after - before) > 1e-6
    # New ALPHA EV/EBITDA: (60*100 + 400 + 10 - 100) / 300 = 21.03x -> median becomes DELTA's 14.0x
    assert abs(after - 14.0) < 1e-6
    assert abs(after_ws.cell(row=_company_rows(after_ws)["ALPHA"], column=cols["EV / EBITDA (TTM)"]).value
               - 6310 / 300) < 1e-6


def test_recalc_bank_workbook(bank_path):
    wb = _recalc_or_skip(bank_path)
    _assert_no_excel_errors(wb)
    ws = wb[SHEET_COMPS]
    cols = _header_cols(ws)
    rows = _company_rows(ws)
    stats = _stat_rows(ws)
    banks = {c.ticker: c for c in bank_companies()}
    for ticker, r in rows.items():
        assert abs(ws.cell(row=r, column=cols["P / TBV"]).value - banks[ticker].multiples["p_tbv"]) < 1e-6
        assert abs(ws.cell(row=r, column=cols["P / E (TTM)"]).value - banks[ticker].multiples["pe_ttm"]) < 1e-6
    ptbv = [float(c.multiples["p_tbv"]) for c in banks.values()]
    assert abs(ws.cell(row=stats["Median"], column=cols["P / TBV"]).value - float(np.median(ptbv))) < 1e-6
