"""
Tests for compsai/pipeline.py: peer-set config, the offline end-to-end run, progress
reporting, per-company failure handling, the bank branch, and a Streamlit smoke test.

Everything runs against the bundled FIXA/FIXB/FIXC fixtures (conftest.py forces offline mode).
"""

from __future__ import annotations

import json
import textwrap
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

from compsai import CONFIG_DIR, PROJECT_ROOT
from compsai import pipeline
from compsai.models import CommentaryResult, PipelineResult
from compsai.valuation import MULTIPLES_BY_SECTOR, MULTIPLES_INDEX, STAT_ROWS, summary_stats

FIXTURE_TICKERS = ["FIXA", "FIXB", "FIXC"]


# ------------------------------------------------------------------------------------------
# Peer sets
# ------------------------------------------------------------------------------------------

def test_load_peer_sets_reads_the_shipped_config():
    sets = pipeline.load_peer_sets(CONFIG_DIR / "peer_sets.yaml")
    assert {"us_large_software", "us_banks", "canadian_cross_listed", "consumer_staples"} <= set(sets)
    assert sets["us_banks"].sector_type == "bank"
    assert sets["us_large_software"].sector_type == "industrial"
    assert sets["us_large_software"].tickers == ["MSFT", "ORCL", "CRM", "ADBE", "NOW", "INTU"]
    assert sets["us_banks"].tickers == ["JPM", "BAC", "WFC", "C", "GS", "MS"]


def test_load_peer_sets_accepts_bare_lists_and_mappings(tmp_path):
    cfg = tmp_path / "peers.yaml"
    cfg.write_text(textwrap.dedent("""
        bare_list: [aapl, msft]
        mapping:
          sector_type: bank
          tickers: [jpm, " bac "]
    """))
    sets = pipeline.load_peer_sets(cfg)
    assert sets["bare_list"].tickers == ["AAPL", "MSFT"]
    assert sets["bare_list"].sector_type == "industrial"
    assert sets["mapping"].tickers == ["JPM", "BAC"]
    assert sets["mapping"].sector_type == "bank"


def test_load_peer_sets_rejects_unknown_sector(tmp_path):
    cfg = tmp_path / "peers.yaml"
    cfg.write_text("oops:\n  sector_type: crypto\n  tickers: [X]\n")
    with pytest.raises(ValueError, match="sector_type"):
        pipeline.load_peer_sets(cfg)


# ------------------------------------------------------------------------------------------
# End to end (offline)
# ------------------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def e2e(tmp_path_factory) -> tuple[PipelineResult, list[tuple[str, str]]]:
    out = tmp_path_factory.mktemp("e2e")
    events: list[tuple[str, str]] = []
    result = pipeline.run_pipeline(
        ["FIXB", "FIXC"], peer_set_name="demo", target="FIXA", sector_type="industrial",
        with_commentary=True, progress=lambda s, m: events.append((s, m)),
        offline=True, out_dir=out, write_excel=True,
    )
    return result, events


def test_e2e_comps_table_shape(e2e):
    result, _ = e2e
    comps = result.comps
    assert list(comps.index) == ["FIXA", "FIXB", "FIXC"]  # target first
    assert comps.index.name == "ticker"
    assert list(comps.columns) == ["name", "is_target"] + MULTIPLES_INDEX
    assert comps.loc["FIXA", "is_target"] and not comps.loc["FIXB", "is_target"]
    assert comps.loc["FIXA", "name"] == "Fixture Fruit Inc."
    # FIXA market fixture: price 200 x 15,000mm shares = $3,000,000mm market cap.
    assert comps.loc["FIXA", "market_cap"] == pytest.approx(3_000_000)
    assert comps["ev_ebitda_ttm"].dtype == "float64"
    assert result.target == "FIXA" and result.sector_type == "industrial" and result.peer_set == "demo"


def test_e2e_stats_exclude_target(e2e):
    result, _ = e2e
    stats = result.stats
    assert list(stats.index) == STAT_ROWS
    assert list(stats.columns) == MULTIPLES_BY_SECTOR["industrial"]
    # Recompute by hand from the peers only (FIXB and FIXC); the target must not be inside.
    expected = summary_stats(result.comps.drop(index="FIXA"), MULTIPLES_BY_SECTOR["industrial"])
    pd.testing.assert_frame_equal(stats, expected)
    # FIXC is a bank fixture with no EBITDA, so the EV/EBITDA statistics come from FIXB alone
    # and mean == median == p25 == p75 == FIXB's multiple.
    fixb = result.comps.loc["FIXB", "ev_ebitda_ttm"]
    assert stats.loc["median", "ev_ebitda_ttm"] == pytest.approx(fixb)
    assert stats.loc["p25", "ev_ebitda_ttm"] == pytest.approx(fixb)


def test_e2e_football_field(e2e):
    result, _ = e2e
    ff = result.football_field
    assert ff is not None
    assert list(ff["method"]) == ["EV / Revenue (TTM)", "EV / EBITDA (TTM)", "P / E (TTM)"]
    assert (ff["implied_price_high"] >= ff["implied_price_low"]).all()


def test_e2e_commentary_from_fixture(e2e):
    result, _ = e2e
    # Every company gets a CommentaryResult: FIXA from its fixture, the others carry the reason
    # the stage produced nothing (so the Commentary sheet says why instead of "not run").
    assert set(result.commentary) == {"FIXA", "FIXB", "FIXC"}
    assert result.commentary["FIXB"].normalization_items == []
    assert result.commentary["FIXB"].errors == ["offline mode: no commentary fixture for this ticker"]
    res = result.commentary["FIXA"]
    assert isinstance(res, CommentaryResult)
    assert len(res.normalization_items) >= 3
    assert all(item["verified"] for item in res.normalization_items)
    assert res.premium_discount["premium_or_discount"] == "premium"
    assert any("FIXB" in w and "commentary" in w for w in result.warnings)


def test_e2e_workbook_written(e2e):
    result, _ = e2e
    assert result.xlsx_path is not None and result.xlsx_path.exists()
    assert result.xlsx_path.name.startswith("comps_demo_")
    wb = load_workbook(result.xlsx_path)
    assert wb.sheetnames == ["Comps", "Inputs", "Football Field", "Commentary"]
    comps_ws = wb["Comps"]
    # Every numeric company cell is a formula into Inputs; stats rows are formulas.
    formulas = [c.value for row in comps_ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("Inputs!" in f for f in formulas)
    assert any(f.startswith("=MEDIAN(") for f in formulas)


def test_e2e_progress_stages_in_order(e2e):
    _, events = e2e
    seen = [stage for stage, _ in events]
    order = [s for i, s in enumerate(seen) if i == 0 or s != seen[i - 1]]  # collapse repeats
    assert order == list(pipeline.STAGES)
    assert any("fixture" in m.lower() or "offline" in m.lower() for s, m in events if s == "edgar")


def test_e2e_warnings_carry_fixture_ttm_notes(e2e):
    result, _ = e2e
    # FIXB's latest filing is its 10-K, so edgar.py reports TTM = FY for it.
    assert any(w.startswith("FIXB:") and "TTM = FY" in w for w in result.warnings)


# ------------------------------------------------------------------------------------------
# Failure handling and options
# ------------------------------------------------------------------------------------------

def test_unknown_ticker_is_skipped_with_warning(tmp_path):
    result = pipeline.run_pipeline(["FIXA", "ZZZZ"], offline=True, with_commentary=False,
                                   write_excel=False)
    assert list(result.comps.index) == ["FIXA"]
    assert any(w.startswith("ZZZZ: skipped") for w in result.warnings)
    assert result.xlsx_path is None and result.football_field is None


def test_all_tickers_failing_raises():
    with pytest.raises(RuntimeError, match="no company could be loaded"):
        pipeline.run_pipeline(["ZZZZ", "YYYY"], offline=True, with_commentary=False, write_excel=False)


def test_missing_target_drops_football_field_but_keeps_peers():
    result = pipeline.run_pipeline(["FIXA"], target="ZZZZ", offline=True, with_commentary=False,
                                   write_excel=False)
    assert result.target is None and result.football_field is None
    assert any("ZZZZ" in w and "target" in w for w in result.warnings)


def test_bad_inputs_raise():
    with pytest.raises(ValueError):
        pipeline.run_pipeline([], offline=True)
    with pytest.raises(ValueError, match="sector_type"):
        pipeline.run_pipeline(["FIXA"], sector_type="crypto", offline=True)


def test_price_overrides_win(tmp_path):
    result = pipeline.run_pipeline(
        ["FIXA"], offline=True, with_commentary=False, write_excel=False,
        price_overrides={"fixa": {"price": 100.0, "shares_outstanding": 10_000.0}},
    )
    assert result.companies[0].market["source"] == "override"
    assert result.comps.loc["FIXA", "market_cap"] == pytest.approx(1_000_000)


def test_bank_sector_branch(tmp_path):
    result = pipeline.run_pipeline(
        ["FIXA", "FIXB"], target="FIXC", sector_type="bank", peer_set_name="demo_bank",
        offline=True, with_commentary=False, out_dir=tmp_path, write_excel=True,
    )
    assert list(result.stats.columns) == ["pe_ttm", "p_tbv"]
    assert result.comps["ev"].isna().all()  # EV is not meaningful for banks
    assert list(result.football_field["method"]) == ["P / E (TTM)", "P / TBV"]
    wb = load_workbook(result.xlsx_path)
    headers = [c.value for c in wb["Comps"][4]]
    assert "P / TBV" in headers and "EV / EBITDA (TTM)" not in headers


def test_commentary_skipped_without_api_key(monkeypatch):
    """Online-style run without a key: the stage reports the skip instead of calling Anthropic."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setattr(pipeline, "load_dotenv", lambda *a, **k: None)
    warnings: list[str] = []
    results = pipeline._run_commentary([], pd.DataFrame(), "industrial", offline=False, refresh=False,
                                       client=None, warnings=warnings, progress=None)
    assert results == {}
    assert warnings == ["ANTHROPIC_API_KEY not set; commentary skipped"]


class _FakeClient:
    """Minimal stand-in for anthropic.Anthropic: returns canned JSON for every call."""

    def __init__(self):
        self.calls: list[dict] = []
        self.messages = self

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if "normalising earnings" in kwargs["messages"][0]["content"]:
            text = json.dumps({"items": [{
                "description": "Orchard Networks patent litigation settlement", "amount_usd_m": 410,
                "fiscal_year": 2024, "direction": "add_back",
                "source_quote": "one-time charge of $410 million to settle the Orchard Networks patent litigation",
            }]})
        else:
            text = json.dumps({"growth_outlook": "solid", "margin_trajectory": "expanding",
                               "key_risks": ["a", "b", "c"], "premium_or_discount": "premium",
                               "rationale": "trades at a premium that is justified"})
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=text)])


def test_injected_commentary_client_runs_the_real_commentary_path():
    fake = _FakeClient()
    result = pipeline.run_pipeline(["FIXA"], offline=True, with_commentary=True, write_excel=False,
                                   commentary_client=fake)
    res = result.commentary["FIXA"]
    assert fake.calls, "the fake client should have been called"
    assert res.premium_discount["premium_or_discount"] == "premium"
    assert res.normalization_items and res.normalization_items[0]["verified"] is True
    assert res.source_url.endswith("/fixa-20240928.htm")


def test_peer_median_covers_margins_and_masks_target():
    comps = pd.DataFrame({
        "name": ["T", "A", "B"], "is_target": [True, False, False],
        "ev_ebitda_ttm": [50.0, 10.0, -5.0],  # -5 is NM and must be masked
        "net_margin": [0.9, 0.1, 0.3],
        "note": ["", "", ""],
    }, index=["T", "A", "B"])
    median = pipeline._peer_median(comps)
    assert median["ev_ebitda_ttm"] == pytest.approx(10.0)
    assert median["net_margin"] == pytest.approx(0.2)


# ------------------------------------------------------------------------------------------
# CLI
# ------------------------------------------------------------------------------------------

def test_cli_offline_run(tmp_path, capsys):
    rc = pipeline.main(["--tickers", "FIXA,FIXB,FIXC", "--target", "FIXA", "--offline",
                        "--out", str(tmp_path), "--peer-set", "us_large_software"])
    assert rc == 0
    out = capsys.readouterr().out
    assert "=== Comps: us_large_software" in out
    assert "Football field: FIXA" in out
    assert "Workbook:" in out
    assert list(tmp_path.glob("comps_us_large_software_*.xlsx"))


def test_cli_name_flag_labels_a_custom_run(tmp_path, capsys):
    rc = pipeline.main(["--tickers", "FIXC,FIXA", "--target", "FIXC", "--sector-type", "bank", "--name", "demo_bank",
                        "--offline", "--no-ai", "--out", str(tmp_path)])
    assert rc == 0
    assert "=== Comps: demo_bank (bank)" in capsys.readouterr().out
    assert list(tmp_path.glob("comps_demo_bank_*.xlsx"))


def test_cli_requires_tickers_or_peer_set():
    with pytest.raises(SystemExit):
        pipeline.main(["--offline"])


# ------------------------------------------------------------------------------------------
# Streamlit app smoke test (skipped when streamlit's test harness is unavailable)
# ------------------------------------------------------------------------------------------

def test_streamlit_app_parses():
    import ast

    source = (PROJECT_ROOT / "app" / "streamlit_app.py").read_text(encoding="utf-8")
    ast.parse(source)


def test_streamlit_app_renders_without_exceptions():
    apptest = pytest.importorskip("streamlit.testing.v1")
    at = apptest.AppTest.from_file(str(PROJECT_ROOT / "app" / "streamlit_app.py"), default_timeout=120)
    at.run()
    assert not at.exception, [str(e) for e in at.exception]
    assert at.sidebar.selectbox, "the peer-set selector should render"


def test_streamlit_app_offline_run_produces_results():
    apptest = pytest.importorskip("streamlit.testing.v1")
    at = apptest.AppTest.from_file(str(PROJECT_ROOT / "app" / "streamlit_app.py"), default_timeout=240)
    at.run()
    assert not at.exception
    # Point the app at the bundled fixtures and run.
    at.sidebar.checkbox(key="offline").set_value(True)
    at.sidebar.text_input(key="custom_tickers").set_value("FIXA, FIXB, FIXC")
    at.sidebar.text_input(key="target").set_value("FIXA")
    at.sidebar.checkbox(key="with_ai").set_value(False)
    at.sidebar.button(key="run").click().run()
    assert not at.exception, [str(e) for e in at.exception]
    assert at.session_state["result"] is not None
    assert list(at.session_state["result"].comps.index) == ["FIXA", "FIXB", "FIXC"]


def test_streamlit_sector_type_follows_the_chosen_peer_set():
    """A keyed selectbox keeps its own state, so picking a bank peer set must push the bank
    branch into session state (otherwise banks would be valued on EV/EBITDA)."""
    apptest = pytest.importorskip("streamlit.testing.v1")
    at = apptest.AppTest.from_file(str(PROJECT_ROOT / "app" / "streamlit_app.py"), default_timeout=120)
    at.run()
    at.sidebar.selectbox(key="peer_set").set_value("us_banks").run()
    assert at.session_state["sector_type"] == "bank"
    at.sidebar.selectbox(key="peer_set").set_value("us_large_software").run()
    assert at.session_state["sector_type"] == "industrial"
